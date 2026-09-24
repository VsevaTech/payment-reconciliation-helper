"""End-to-end through FastAPI: upload → map → reconcile → export, plus JSON API."""

from __future__ import annotations

import csv
import io
import json
import re
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.models import RECONCILED_CATEGORIES

ROOT = Path(__file__).resolve().parents[1]
EXPECTED = json.loads((ROOT / "demo-data" / "expected.json").read_text(encoding="utf-8"))


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def demo_files() -> dict:
    return {
        "orders": ("orders.csv", (ROOT / "demo-data" / "orders.csv").read_bytes(), "text/csv"),
        "payments": (
            "payments.csv",
            (ROOT / "demo-data" / "payments.csv").read_bytes(),
            "text/csv",
        ),
    }


def test_health(client: TestClient):
    assert client.get("/health").json()["status"] == "ok"


def test_index_renders(client: TestClient):
    r = client.get("/")
    assert r.status_code == 200 and "Upload files" in r.text


def test_full_ui_flow(client: TestClient):
    r = client.post("/upload", files=demo_files(), follow_redirects=False)
    assert r.status_code == 303
    map_url = r.headers["location"]
    session_id = re.search(r"/s/([0-9a-f]+)/map", map_url).group(1)

    page = client.get(map_url)
    assert page.status_code == 200
    # auto-guessed mapping is pre-selected
    assert '<option value="Order ID" selected>' in page.text
    assert '<option value="merchant_reference" selected>' in page.text
    assert "<code>utf-8</code>" in page.text
    assert "<code>;</code>" in page.text

    form = {
        "orders_reference": "Order ID",
        "orders_amount": "Amount",
        "orders_currency": "Currency",
        "orders_date": "Date",
        "payments_reference": "merchant_reference",
        "payments_amount": "amount",
        "payments_currency": "currency",
        "payments_date": "payment_date",
        "payments_transaction_type": "transaction_type",
        "tolerance": "0",
        "window_days": "3",
    }
    r = client.post(f"/s/{session_id}/reconcile", data=form, follow_redirects=False)
    assert r.status_code == 303

    result = client.get(f"/s/{session_id}/result")
    assert result.status_code == 200
    for cat, count in EXPECTED["summary"].items():
        tile = f'<div class="n">{count}</div><div class="k"><span class="pill {cat}">'
        assert tile in result.text

    partial = client.get(
        f"/s/{session_id}/result?category=AMOUNT_MISMATCH", headers={"HX-Request": "true"}
    )
    assert "<html" not in partial.text
    assert partial.text.count('class="pill AMOUNT_MISMATCH"') == 3

    csv_resp = client.get(f"/s/{session_id}/export.csv")
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-disposition"].endswith('filename="reconciliation.csv"')
    reconciled = {c.value for c in RECONCILED_CATEGORIES}
    exceptions_total = sum(v for k, v in EXPECTED["summary"].items() if k not in reconciled)
    assert len(csv_resp.content.decode("utf-8-sig").strip().splitlines()) == exceptions_total + 1

    xlsx_resp = client.get(f"/s/{session_id}/export.xlsx")
    assert xlsx_resp.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx_resp.content))
    assert wb["Exceptions"].max_row == exceptions_total + 1
    assert wb["All rows"].max_row == sum(EXPECTED["summary"].values()) + 1


def test_upload_rejects_invalid_csv(client: TestClient):
    files = demo_files()
    files["payments"] = ("payments.csv", b"\x00\x01\x02\x03garbage", "text/csv")
    r = client.post("/upload", files=files)
    assert r.status_code == 200
    assert "payments:" in r.text and "binary" in r.text


def test_reconcile_with_bad_mapping_shows_error(client: TestClient):
    r = client.post("/upload", files=demo_files(), follow_redirects=False)
    session_id = re.search(r"/s/([0-9a-f]+)/map", r.headers["location"]).group(1)
    r = client.post(
        f"/s/{session_id}/reconcile",
        data={
            "orders_reference": "Nope",
            "orders_amount": "Amount",
            "payments_reference": "merchant_reference",
            "payments_amount": "amount",
        },
    )
    assert r.status_code == 200 and "does not exist" in r.text


def test_unknown_session_404(client: TestClient):
    assert client.get("/s/deadbeef/result").status_code == 404


def test_json_api(client: TestClient):
    r = client.post(
        "/api/reconcile",
        files=demo_files(),
        data={
            "orders_reference": "Order ID",
            "orders_amount": "Amount",
            "orders_currency": "Currency",
            "orders_date": "Date",
            "payments_reference": "merchant_reference",
            "payments_amount": "amount",
            "payments_currency": "currency",
            "payments_date": "payment_date",
            "payments_transaction_type": "transaction_type",
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"] == EXPECTED["summary"]
    assert body["financials"] == EXPECTED["financials"]
    assert body["transaction_type_mapped"] is True
    assert body["detected"]["payments"]["delimiter"] == ";"
    assert len(body["rows"]) == sum(EXPECTED["summary"].values())
    assert all(isinstance(row["order_amount"], str) for row in body["rows"])  # never float


def test_json_api_invalid_tolerance(client: TestClient):
    r = client.post(
        "/api/reconcile",
        files=demo_files(),
        data={
            "orders_reference": "Order ID",
            "orders_amount": "Amount",
            "payments_reference": "merchant_reference",
            "payments_amount": "amount",
            "tolerance": "lots",
        },
    )
    assert r.status_code == 400


def _typed_session(client: TestClient) -> str:
    r = client.post("/upload", files=demo_files(), follow_redirects=False)
    session_id = re.search(r"/s/([0-9a-f]+)/map", r.headers["location"]).group(1)
    page = client.get(f"/s/{session_id}/map")
    # transaction_type is guessed from the header, "status" is not
    assert '<option value="transaction_type" selected>' in page.text
    form = {
        "orders_reference": "Order ID",
        "orders_amount": "Amount",
        "orders_currency": "Currency",
        "orders_date": "Date",
        "payments_reference": "merchant_reference",
        "payments_amount": "amount",
        "payments_currency": "currency",
        "payments_date": "payment_date",
        "payments_transaction_type": "transaction_type",
    }
    client.post(f"/s/{session_id}/reconcile", data=form, follow_redirects=False)
    return session_id


def test_result_page_financial_summary_per_currency(client: TestClient):
    sid = _typed_session(client)
    page = client.get(f"/s/{sid}/result").text
    assert "Financial summary by currency" in page
    for t in EXPECTED["financials"]:
        assert f"<strong>{t['currency']}</strong>" in page
    aed = next(t for t in EXPECTED["financials"] if t["currency"] == "AED")
    net = f"{Decimal(aed['net_captured']):,.2f}"
    assert f"<strong>{net}</strong>" in page
    assert "not mapped" not in page.split("Financial summary")[0]
    assert "Charged back" in page and "CB reversed" in page
    assert f'<td class="num">{Decimal(aed["charged_back"]):,.2f}</td>' in page


@pytest.mark.parametrize(
    ("flt", "categories"),
    [
        ("split", ["MATCHED_SPLIT"]),
        ("partial", ["PARTIAL_PAYMENT"]),
        ("refunded", ["PARTIALLY_REFUNDED", "REFUNDED"]),
        ("missing", ["MISSING_PAYMENT"]),
        ("amount_mismatch", ["AMOUNT_MISMATCH"]),
        ("duplicates", ["POSSIBLE_DUPLICATE_CAPTURE", "DUPLICATE_ORDER"]),
        ("orphans", ["ORPHAN_PAYMENT"]),
        ("matched", ["MATCHED"]),
        (
            "chargebacks",
            [
                "CHARGED_BACK",
                "CHARGEBACK_REVERSED",
                "CHARGEBACK_EXCEEDS_CAPTURE",
                "REVERSAL_EXCEEDS_CHARGEBACK",
            ],
        ),
    ],
)
def test_result_filters(client: TestClient, flt: str, categories: list[str]):
    sid = _typed_session(client)
    partial = client.get(f"/s/{sid}/result?filter={flt}", headers={"HX-Request": "true"}).text
    assert "<html" not in partial
    want = sum(EXPECTED["summary"][c] for c in categories)
    got = sum(partial.count(f'<span class="pill {c}">') for c in categories)
    assert got == min(want, 2000)
    assert f"— {want} rows" in partial


def test_result_currency_filter_and_unknown_filter(client: TestClient):
    sid = _typed_session(client)
    page = client.get(f"/s/{sid}/result?filter=all&currency=EUR").text
    eur = next(t for t in EXPECTED["financials"] if t["currency"] == "EUR")
    assert "All rows · EUR" in page
    # at least every EUR order is listed
    assert page.count("<td>EUR") >= eur["orders_count"]
    assert client.get(f"/s/{sid}/result?filter=bogus").status_code == 400


def test_financials_csv_and_xlsx_export(client: TestClient):
    sid = _typed_session(client)
    r = client.get(f"/s/{sid}/export-financials.csv")
    assert r.status_code == 200
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [{k: v for k, v in row.items()} for row in rows] == [
        {k: str(v) for k, v in t.items()} for t in EXPECTED["financials"]
    ]
    wb = load_workbook(io.BytesIO(client.get(f"/s/{sid}/export.xlsx").content))
    values = [row[0] for row in wb["Summary"].iter_rows(values_only=True)]
    assert "Financial summary by currency" in values
    all_rows = wb["All rows"]
    header = [c.value for c in all_rows[1]]
    assert "refunded_amount" in header and "payment_rows" in header


def test_json_api_without_transaction_type_keeps_legacy_mode(client: TestClient):
    r = client.post(
        "/api/reconcile",
        files=demo_files(),
        data={
            "orders_reference": "Order ID",
            "orders_amount": "Amount",
            "orders_currency": "Currency",
            "payments_reference": "merchant_reference",
            "payments_amount": "amount",
            "payments_currency": "currency",
            "payments_date": "payment_date",
            "include_rows": "false",
        },
    )
    body = r.json()
    assert body["transaction_type_mapped"] is False
    assert body["summary"] == EXPECTED["summary_without_transaction_type"]
    assert "rows" not in body


def test_json_api_rejects_unknown_transaction_type_column(client: TestClient):
    r = client.post(
        "/api/reconcile",
        files=demo_files(),
        data={
            "orders_reference": "Order ID",
            "orders_amount": "Amount",
            "payments_reference": "merchant_reference",
            "payments_amount": "amount",
            "payments_transaction_type": "nope",
        },
    )
    assert r.status_code == 400 and "does not exist" in r.text


def test_cli_does_not_adopt_a_guessed_type_column_with_foreign_values(tmp_path: Path):
    import subprocess
    import sys

    (tmp_path / "o.csv").write_text("order_id,amount,currency\nA,100,USD\n")
    (tmp_path / "p.csv").write_text("merchant_reference,amount,currency,type\nA,100,USD,card\n")
    proc = subprocess.run(
        [
            sys.executable,
            "-m",
            "app.cli",
            str(tmp_path / "o.csv"),
            str(tmp_path / "p.csv"),
            "--json",
        ],
        capture_output=True,
        text=True,
        cwd=ROOT,
        check=True,
    )
    assert json.loads(proc.stdout)["MATCHED"] == 1
    assert "was not used as transaction type" in proc.stderr


def test_map_page_does_not_preselect_implausible_type_column(client: TestClient):
    files = {
        "orders": ("o.csv", b"order_id,amount\nA,1\n", "text/csv"),
        "payments": ("p.csv", b"merchant_reference,amount,type\nA,1,card\n", "text/csv"),
    }
    r = client.post("/upload", files=files, follow_redirects=False)
    page = client.get(r.headers["location"]).text
    assert '<option value="type" selected>' not in page
    assert '<option value="type" >' in page or '<option value="type">' in page


def test_empty_currency_can_be_filtered(client: TestClient):
    files = {
        "orders": ("o.csv", b"order_id,amount,currency\nA,100,\nB,5,USD\n", "text/csv"),
        "payments": ("p.csv", b"merchant_reference,amount,currency\nA,100,\nB,5,USD\n", "text/csv"),
    }
    r = client.post("/upload", files=files, follow_redirects=False)
    sid = re.search(r"/s/([0-9a-f]+)/map", r.headers["location"]).group(1)
    form = {
        "orders_reference": "order_id",
        "orders_amount": "amount",
        "orders_currency": "currency",
        "payments_reference": "merchant_reference",
        "payments_amount": "amount",
        "payments_currency": "currency",
    }
    client.post(f"/s/{sid}/reconcile", data=form, follow_redirects=False)
    page = client.get(f"/s/{sid}/result?filter=all&currency=-").text
    assert "All rows · no currency" in page and "— 1 rows" in page
    assert "?filter=all&currency=-" in client.get(f"/s/{sid}/result").text
