from __future__ import annotations

import csv
import io
from decimal import Decimal

from openpyxl import load_workbook

from app.export import COLUMNS, FINANCIAL_COLUMNS, to_csv, to_xlsx
from app.models import ColumnMapping
from app.reconcile import reconcile_frames
from tests.conftest import ORDERS_MAP, frame, run

ORDERS = "order_id,amount,currency,date\nA-1,100.00,AED,2024-09-01\nA-2,0.10,AED,2024-09-01\n"
PAYMENTS = (
    "merchant_reference,amount,currency,payment_date\n"
    "A-1,100.00,AED,2024-09-01\nA-2,0.30,AED,2024-09-01\n"
)


def test_csv_export_has_bom_header_and_exact_decimals():
    result = run(ORDERS, PAYMENTS)
    data = to_csv(result, include_matched=True)
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert list(rows[0].keys()) == COLUMNS
    mismatch = next(r for r in rows if r["category"] == "AMOUNT_MISMATCH")
    assert mismatch["difference"] == "0.20"
    only_exc = list(csv.DictReader(io.StringIO(to_csv(result, False).decode("utf-8-sig"))))
    assert [r["category"] for r in only_exc] == ["AMOUNT_MISMATCH"]


def test_xlsx_export_sheets_and_numeric_cells():
    result = run(ORDERS, PAYMENTS)
    wb = load_workbook(io.BytesIO(to_xlsx(result)))
    assert wb.sheetnames == ["Summary", "Exceptions", "All rows"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=9, values_only=True)}
    assert summary["MATCHED"] == 1 and summary["AMOUNT_MISMATCH"] == 1
    exc = wb["Exceptions"]
    header = [c.value for c in exc[1]]
    assert header == COLUMNS
    row = [c.value for c in exc[2]]
    assert row[0] == "AMOUNT_MISMATCH"
    assert Decimal(str(row[COLUMNS.index("difference")])) == Decimal("0.2")
    assert exc.freeze_panes == "A2"


TYPED_ORDERS = (
    "order_id,amount,currency,date\nR-1,100.00,AED,2024-09-01\nS-1,100.00,USD,2024-09-01\n"
)
TYPED_PAYMENTS = (
    "merchant_reference,amount,currency,payment_date,type\n"
    "R-1,100.00,AED,2024-09-01,CAPTURE\n"
    "S-1,60.00,USD,2024-09-01,CAPTURE\n"
    "R-1,-30.00,AED,2024-09-04,REFUND\n"
    "S-1,40.00,USD,2024-09-02,CAPTURE\n"
)


def _typed_result():
    return reconcile_frames(
        frame(TYPED_ORDERS),
        frame(TYPED_PAYMENTS),
        ORDERS_MAP,
        ColumnMapping(
            reference="merchant_reference",
            amount="amount",
            currency="currency",
            date="payment_date",
            transaction_type="type",
        ),
    )


def test_csv_export_carries_refund_split_columns_and_source_rows():
    rows = list(csv.DictReader(io.StringIO(to_csv(_typed_result()).decode("utf-8-sig"))))
    by_ref = {r["reference"]: r for r in rows}
    refund = by_ref["R-1"]
    assert refund["category"] == "PARTIALLY_REFUNDED"
    assert (refund["captured_amount"], refund["refunded_amount"], refund["net_amount"]) == (
        "100.00",
        "30.00",
        "70.00",
    )
    assert refund["payment_rows"] == "1;3" and refund["transaction_types"] == "PAYMENT, REFUND"
    split = by_ref["S-1"]
    assert split["category"] == "MATCHED_SPLIT" and split["payment_rows"] == "2;4"
    assert split["payment_amount"] == "100.00"  # backwards-compatible: sum of captures


def test_reconciled_categories_are_not_exceptions_in_export():
    only_exc = to_csv(_typed_result(), include_matched=False).decode("utf-8-sig")
    assert only_exc.strip() == ",".join(COLUMNS)


def test_xlsx_financial_summary_never_mixes_currencies():
    wb = load_workbook(io.BytesIO(to_xlsx(_typed_result())))
    rows = list(wb["Summary"].iter_rows(values_only=True))
    header_at = next(i for i, r in enumerate(rows) if r[0] == "currency")
    assert list(rows[header_at])[: len(FINANCIAL_COLUMNS)] == FINANCIAL_COLUMNS
    fin = {r[0]: r for r in rows[header_at + 1 :] if r[0]}
    assert set(fin) == {"AED", "USD"}
    col = FINANCIAL_COLUMNS.index
    assert Decimal(str(fin["AED"][col("net_captured")])) == Decimal("70")
    assert Decimal(str(fin["USD"][col("captured")])) == Decimal("100")
    all_rows = wb["All rows"]
    header = [c.value for c in all_rows[1]]
    r1 = next(r for r in all_rows.iter_rows(min_row=2, values_only=True) if r[1] == "R-1")
    assert Decimal(str(r1[header.index("refunded_amount")])) == Decimal("30")
    assert r1[header.index("payment_rows")] == "1;3"
