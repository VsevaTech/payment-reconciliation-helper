"""CSV and XLSX report writers."""

from __future__ import annotations

import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import CATEGORY_ORDER, Category, ReconcileResult, ResultRow

COLUMNS = [
    "category",
    "reference",
    "order_row",
    "payment_row",
    "payment_rows",
    "order_amount",
    "payment_amount",
    "captured_amount",
    "refunded_amount",
    "chargeback_amount",
    "chargeback_reversed_amount",
    "voided_amount",
    "net_amount",
    "difference",
    "order_currency",
    "payment_currency",
    "transaction_types",
    "duplicate_payment_rows",
    "order_date",
    "payment_date",
    "explanation",
]

MONEY_COLUMNS = (
    "order_amount",
    "payment_amount",
    "captured_amount",
    "refunded_amount",
    "chargeback_amount",
    "chargeback_reversed_amount",
    "voided_amount",
    "net_amount",
    "difference",
)

FINANCIAL_COLUMNS = [
    "currency",
    "orders_count",
    "orders_total",
    "captured",
    "refunded",
    "charged_back",
    "chargeback_reversed",
    "net_captured",
    "voided",
    "unreconciled",
]

_FILL = {
    Category.MATCHED: "E6F4EA",
    Category.MATCHED_SPLIT: "E6F4EA",
    Category.PARTIALLY_REFUNDED: "E0F0FF",
    Category.REFUNDED: "E0F0FF",
    Category.CHARGEBACK_REVERSED: "E0F0FF",
    Category.MISSING_PAYMENT: "FDE7E9",
    Category.ORPHAN_PAYMENT: "FDE7E9",
    Category.AMOUNT_MISMATCH: "FFF4CE",
    Category.PARTIAL_PAYMENT: "FFF4CE",
    Category.CURRENCY_MISMATCH: "FFF4CE",
    Category.POSSIBLE_DUPLICATE_CAPTURE: "E8E0F7",
    Category.REFUND_EXCEEDS_CAPTURE: "FDE7E9",
    Category.CHARGED_BACK: "FFE3CC",
    Category.CHARGEBACK_EXCEEDS_CAPTURE: "FDE7E9",
    Category.REVERSAL_EXCEEDS_CHARGEBACK: "FDE7E9",
    Category.VOIDED: "EEEEEE",
    Category.DUPLICATE_ORDER: "E8E0F7",
    Category.FALLBACK_MATCHED: "E0F0FF",
    Category.INVALID_ROW: "EEEEEE",
}


def _rows_for_export(result: ReconcileResult, include_matched: bool) -> list[ResultRow]:
    return result.rows if include_matched else result.exceptions


def financials_to_csv(result: ReconcileResult) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=FINANCIAL_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for t in result.financials:
        writer.writerow(t.to_dict())
    return buf.getvalue().encode("utf-8-sig")


def to_csv(result: ReconcileResult, include_matched: bool = True) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=COLUMNS, lineterminator="\n")
    writer.writeheader()
    for r in _rows_for_export(result, include_matched):
        writer.writerow(r.to_dict())
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(result: ReconcileResult, include_matched: bool = True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)

    ws.append(["Payment Reconciliation Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Orders rows", result.orders_total])
    ws.append(["Payments rows", result.payments_total])
    ws.append(["Amount tolerance", format(result.options.amount_tolerance, "f")])
    ws.append(["Fallback matching", "on" if result.options.enable_fallback_matching else "off"])
    ws.append([])
    ws.append(["Category", "Count"])
    ws["A8"].font = bold
    ws["B8"].font = bold
    summary = result.summary
    for cat in CATEGORY_ORDER:
        ws.append([cat.value, summary[cat.value]])

    # Financial totals, one row per currency — currencies are never summed.
    ws.append([])
    ws.append(["Financial summary by currency"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append(
        [
            "Transaction type column",
            "mapped"
            if result.transaction_type_mapped
            else "not mapped (refunds / chargebacks not identified)",
        ]
    )
    ws.append(FINANCIAL_COLUMNS)
    for cell in ws[ws.max_row]:
        cell.font = bold
    for t in result.financials:
        d = t.to_dict()
        ws.append(
            [
                d[c] if c in ("currency", "orders_count") else Decimal(str(d[c]))
                for c in FINANCIAL_COLUMNS
            ]
        )
        for col in range(3, len(FINANCIAL_COLUMNS) + 1):
            ws.cell(row=ws.max_row, column=col).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 28
    for letter in "BCDEFGHIJ":
        ws.column_dimensions[letter].width = 16

    def write_sheet(title: str, rows: list[ResultRow]) -> None:
        sheet = wb.create_sheet(title)
        sheet.append(COLUMNS)
        for cell in sheet[1]:
            cell.font = bold
        for r in rows:
            d = r.to_dict()
            values: list[object] = []
            for col in COLUMNS:
                v = d[col]
                if col in MONEY_COLUMNS and v != "":
                    # openpyxl writes Decimal as an exact numeric cell.
                    v = Decimal(str(v))
                values.append(v)
            sheet.append(values)
            fill = _FILL.get(r.category)
            if fill:
                sheet.cell(row=sheet.max_row, column=1).fill = PatternFill("solid", fgColor=fill)
        widths = {"category": 28, "reference": 22, "explanation": 90, "transaction_types": 20}
        for i, col in enumerate(COLUMNS, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
        for i, col in enumerate(COLUMNS, start=1):
            if col not in MONEY_COLUMNS:
                continue
            for cells in sheet.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cells:
                    c.number_format = "#,##0.00"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    write_sheet("Exceptions", result.exceptions)
    if include_matched:
        write_sheet("All rows", result.rows)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
